import os
import time
from datetime import datetime
from watchdog.observers import Observer
from watchdog.events import FileSystemEventHandler
import requests
from docx import Document as DocxDocument
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from app.common.path import INPUT_DIR, OUTPUT_DIR

# 配置
API_URL = "http://127.0.0.1:8001/api/v1/chat/stream"
SUPPORTED_EXTENSIONS = (".txt", ".docx", ".pdf", ".png", ".jpg", ".jpeg")

INPUT_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def save_as_docx(text, filepath, filename):
    doc = DocxDocument()

    # 页面边距
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # 设置默认字体
    style = doc.styles['Normal']
    style.font.name = '宋体'
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    # 标题
    title = doc.add_heading('法律分析报告', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.name = '宋体'
        run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        run.font.size = Pt(22)

    # 元信息
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = meta.add_run(f"原始文件: {filename}")
    r1.font.size = Pt(10)
    r1.font.name = '宋体'
    r1.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    r2 = meta.add_run(f"\n生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    r2.font.size = Pt(10)
    r2.font.name = '宋体'
    r2.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    # 分隔线
    sep = doc.add_paragraph("─" * 50)
    for run in sep.runs:
        run.font.name = '宋体'
        run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    # 正文
    for line in text.split("\n"):
        line = line.strip()
        if not line:
            continue

        if line.startswith("###"):
            h = doc.add_heading(line.lstrip("#").strip(), level=3)
            for run in h.runs:
                run.font.name = '宋体'
                run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        elif line.startswith("##"):
            h = doc.add_heading(line.lstrip("#").strip(), level=2)
            for run in h.runs:
                run.font.name = '宋体'
                run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        elif line.startswith("#"):
            h = doc.add_heading(line.lstrip("#").strip(), level=1)
            for run in h.runs:
                run.font.name = '宋体'
                run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        elif line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(line.lstrip("- *").strip(), style='List Bullet')
            for run in p.runs:
                run.font.name = '宋体'
                run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        else:
            p = doc.add_paragraph(line)
            for run in p.runs:
                run.font.name = '宋体'
                run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    doc.save(filepath)


def read_txt(filepath):
    with open(filepath, "r", encoding="utf-8") as f:
        return f.read()


def read_docx(filepath):
    from docx import Document
    doc = Document(filepath)
    return "\n".join([p.text for p in doc.paragraphs])


def read_pdf(filepath):
    import fitz
    doc = fitz.open(filepath)
    text = ""
    for page in doc:
        text += page.get_text()
    doc.close()
    return text


def read_image(filepath):
    return f"[图片文件]: {filepath}"


READERS = {
    ".txt": read_txt,
    ".docx": read_docx,
    ".pdf": read_pdf,
    ".png": read_image,
    ".jpg": read_image,
    ".jpeg": read_image,
}


class LegalFileHandler(FileSystemEventHandler):
    def on_created(self, event):
        if event.is_directory:
            return
        filepath = event.src_path
        ext = os.path.splitext(filepath)[1].lower()
        if ext not in SUPPORTED_EXTENSIONS:
            return

        filename = os.path.basename(filepath)
        print(f"发现新文件: {filename}")
        time.sleep(1)

        try:
            reader = READERS.get(ext, read_txt)
            content = reader(filepath)

            if ext in (".png", ".jpg", ".jpeg"):
                payload = {
                    "message": "请分析这张图片中的法律文件内容，并生成分析报告。",
                    "image_url": f"file:///{filepath.replace(chr(92), '/')}",
                    "thread_id": f"auto-{datetime.now().timestamp()}"
                }
            else:
                payload = {
                    "message": f"请根据以下材料，生成一份专业的法律分析报告：\n\n{content}",
                    "thread_id": f"auto-{datetime.now().timestamp()}"
                }

            print(f"正在分析: {filename}")
            response = requests.post(API_URL, json=payload, stream=True)

            full_response = ""
            for chunk in response.iter_content(chunk_size=None, decode_unicode=True):
                if chunk:
                    full_response += chunk

            base_name = f"分析报告_{os.path.splitext(filename)[0]}_{datetime.now().strftime('%Y%m%d_%H%M%S')}"

            # Markdown
            md_path = os.path.join(OUTPUT_DIR, base_name + ".md")
            with open(md_path, "w", encoding="utf-8") as f:
                f.write(
                    f"# 法律分析报告\n\n**原始文件**: {filename}\n\n**生成时间**: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n---\n\n{full_response}")
            print(f"Markdown 报告已生成: {base_name}.md")

            # Word 文档
            docx_path = os.path.join(OUTPUT_DIR, base_name + ".docx")
            save_as_docx(full_response, docx_path, filename)
            print(f"Word 报告已生成: {base_name}.docx")

        except Exception as e:
            print(f"处理 {filename} 时出错: {e}")


if __name__ == "__main__":
    print(f"正在监控文件夹: {INPUT_DIR}")
    print(f"支持格式: {SUPPORTED_EXTENSIONS}")
    print(f"报告保存到: {OUTPUT_DIR}")
    print("将文件拖入 input 文件夹即可自动生成报告...")

    event_handler = LegalFileHandler()
    observer = Observer()
    observer.schedule(event_handler, str(INPUT_DIR), recursive=False)
    observer.start()

    try:
        while True:
            time.sleep(1)
    except KeyboardInterrupt:
        observer.stop()
        print("监控已停止")
    observer.join()
