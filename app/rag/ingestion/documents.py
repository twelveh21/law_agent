import os
import re
from functools import lru_cache
from html import unescape
from pathlib import Path
from typing import Any, List, Optional


def _preprocess_markdown_for_embedding(text: str) -> str:
    """清理 Markdown 标记但保留标题、条文和表格文字，供 embedding 使用。"""
    if not isinstance(text, str):
        raise TypeError("待预处理的 Markdown 内容必须是字符串")

    text = unescape(text).replace("\r\n", "\n").replace("\r", "\n")
    # 图片只保留替代文字；链接只保留可读的链接标题，不把 URL 当正文语义。
    text = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
    # PDF 转换时添加的页码标题主要是定位信息，避免它干扰语义向量。
    text = re.sub(r"(?m)^\s*#{1,6}\s*第\s*\d+\s*页\s*$", "", text)
    # 去掉 Markdown 标题符号、代码围栏和水平分隔线，但保留标题/代码本身。
    text = re.sub(r"(?m)^\s*#{1,6}\s+", "", text)
    text = re.sub(r"(?m)^\s*```[^\n]*$", "", text)
    text = re.sub(r"(?m)^\s*(?:---+|___+|\*\*\*+)\s*$", "", text)
    # 删除 Markdown 表格的分隔线，并保留单元格文字、去掉管道分隔符。
    text = re.sub(r"(?m)^\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*$", "", text)
    text = text.replace("|", " ")
    # 去掉粗体、斜体、删除线等 Markdown 装饰符号，不删除其中的文字。
    text = re.sub(r"(\*\*|__|~~)(.*?)\1", r"\2", text, flags=re.DOTALL)
    text = re.sub(r"(?<![\w\\])\*(.+?)\*(?!\w)", r"\1", text, flags=re.DOTALL)
    text = re.sub(r"(?<![\w\\])_(.+?)_(?!\w)", r"\1", text, flags=re.DOTALL)
    # 简单 HTML 标签转空格；标签内部文本仍保留。
    text = re.sub(r"</?[A-Za-z][A-Za-z0-9:-]*(?:\s+[^<>]*)?\s*/?>", " ", text)

    cleaned_lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    text = "\n".join(line for line in cleaned_lines if line)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _enhanced_pdf_processing(path: str) -> str:
    """提取 PDF 每一页的文字，并用 Markdown 页标题保留页码边界。

    这是文本型 PDF 的解析，不包含扫描件 OCR；扫描件通常会得到空文本。
    """
    pdf_path = Path(path)
    if not pdf_path.is_file():
        raise FileNotFoundError(f"PDF 文件不存在: {pdf_path}")
    if pdf_path.suffix.lower() != ".pdf":
        raise ValueError(f"文件不是 PDF: {pdf_path}")

    try:
        import fitz
    except ImportError as exc:
        raise RuntimeError("缺少 PyMuPDF，请安装项目依赖 pymupdf") from exc

    page_sections: List[str] = []
    try:
        # with 结束时会自动关闭 PDF 文件句柄，即使提取过程中发生异常也一样。
        with fitz.open(pdf_path) as document:
            for page_number, page in enumerate(document, start=1):
                page_text = page.get_text("text", sort=True).strip()
                if page_text:
                    page_sections.append(f"## 第 {page_number} 页\n\n{page_text}")
    except Exception as exc:
        raise RuntimeError(f"读取 PDF 失败: {pdf_path}: {exc}") from exc

    if not page_sections:
        print(f"[RAG] PDF 未提取到可用文字，可能是扫描件，需要 OCR: {pdf_path}")
        return ""

    markdown_text = "\n\n".join(page_sections)
    print(f"[RAG] PDF 提取成功: {pdf_path} -> {len(markdown_text)} chars")
    return markdown_text


@lru_cache(maxsize=1)
def _get_markdown_instance() -> Optional[Any]:
    """懒加载并缓存 MarkItDown 转换器；不可用时返回 None 走后备读取。"""
    try:
        from markitdown import MarkItDown

        return MarkItDown()
    except Exception as exc:
        # MarkItDown 是统一格式转换器；初始化失败不应阻断支持的基础文本格式。
        print(f"[RAG] MarkItDown 不可用，将使用后备读取方式: {exc}")
        return None


def _fallback_text_reader(path: str) -> str:
    """MarkItDown 不可用或转换失败时，读取常见文本、HTML 和 DOCX 文件。"""
    file_path = Path(path)
    if not file_path.is_file():
        return ""

    suffix = file_path.suffix.lower()

    # 这些格式本身就是文本，按常见编码依次解码。
    text_extensions = {
        ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".jsonl",
        ".xml", ".yaml", ".yml", ".log", ".py",
    }
    if suffix in text_extensions:
        raw_bytes = file_path.read_bytes()
        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                text = raw_bytes.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            # 对无法识别编码的文本尽量保留可读部分，不因少数字节失败。
            text = raw_bytes.decode("utf-8", errors="replace")
        return text.strip()

    # HTML 需要移除标签，返回正文文字，而不是原始标签源码。
    if suffix in {".html", ".htm"}:
        raw_bytes = file_path.read_bytes()
        html = raw_bytes.decode("utf-8", errors="replace")
        try:
            from bs4 import BeautifulSoup

            return BeautifulSoup(html, "html.parser").get_text("\n", strip=True)
        except ImportError as exc:
            raise RuntimeError("读取 HTML 后备格式需要 beautifulsoup4") from exc

    # DOCX 后备读取：保留段落和表格的基本内容，不尝试还原复杂版式。
    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("读取 DOCX 后备格式需要 python-docx") from exc

        document = Document(str(file_path))
        sections: List[str] = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.replace("\n", " ").strip() for cell in row.cells]
                sections.append(" | ".join(cells))
        return "\n\n".join(sections)

    print(f"[RAG] 后备读取暂不支持此格式: {file_path.suffix or '无扩展名'}")
    return ""


def _convert_to_markdown(path:str)->str:
    """
    将任意格式文档转化成markdown文本
    """
    if not os.path.exists(path):
        return ""

    #对pdf文件增强处理
    ext=(os.path.splitext(path)[1] or '').lower()
    if ext=='.pdf':
        return _enhanced_pdf_processing(path)

    #其他格式用markdown统一转换
    md_instance=_get_markdown_instance()
    if md_instance is None:
        return _fallback_text_reader(path)

    try:
        result = md_instance.convert(path)
        markdown_text=getattr(result,"text_content",None)
        if isinstance(markdown_text,str) and markdown_text.strip():
            print(f"[RAG]MarkItDown转换成功：{path}->{len(markdown_text)}chars Markdown")
            return markdown_text
        print(f"[RAG] MarkItDown 未提取到正文，尝试后备读取: {path}")
        return _fallback_text_reader(path)
    except Exception as e:
        print(f"转换失败{path}:{e}")
        return _fallback_text_reader(path)
